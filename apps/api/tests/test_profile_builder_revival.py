from io import BytesIO
from pathlib import Path
from docx import Document
from docx.shared import Mm
import pytest
from cv_validator.profile_builder import (
    AnonymizationPolicy, ProfileLink, apply_profile_anonymization, default_profile_template,
    render_candidate_profile_docx, ProfileBuilderPreferences,
)
from cv_validator.profile_builder_ai import _materialize_candidate_profile
from test_profile_builder import _client, _Extractor, _extracted_payload, PROFILE_TOKEN


def candidate():
    return _materialize_candidate_profile(_extracted_payload())


def test_profile_exports_real_a4_instead_of_word_default_letter():
    doc = Document(BytesIO(render_candidate_profile_docx(candidate(), AnonymizationPolicy())))
    assert abs(doc.sections[0].page_width - Mm(210)) < 1000
    assert abs(doc.sections[0].page_height - Mm(297)) < 1000


def test_old_blind_settings_hide_other_links_but_explicit_reveal_is_preserved():
    legacy = {"hide_linkedin": True, "hide_github": True, "hide_portfolio": True}
    assert AnonymizationPolicy.model_validate(legacy).hide_other_links is True
    assert ProfileBuilderPreferences().anonymization.hide_other_links is True
    assert AnonymizationPolicy.model_validate({**legacy, "hide_other_links": False}).hide_other_links is False
    assert AnonymizationPolicy().hide_other_links is False


def test_other_links_are_hidden_in_export_not_in_saved_canonical_profile():
    profile = candidate()
    profile.personal.links.other = [ProfileLink(label="Personal site", url="https://example.test/private-candidate")]
    hidden = AnonymizationPolicy(hide_other_links=True)
    assert apply_profile_anonymization(profile, hidden).personal.links.other == []
    assert len(profile.personal.links.other) == 1
    doc = Document(BytesIO(render_candidate_profile_docx(profile, hidden)))
    assert "private-candidate" not in "\n".join(p.text for p in doc.paragraphs)


def test_profile_read_and_export_responses_cannot_be_cached(tmp_path, location_resolver):
    client = _client(tmp_path, location_resolver, _Extractor(_extracted_payload()))
    response = client.get("/profile-builder/profiles", headers={"X-Profile-Builder-Access-Token": PROFILE_TOKEN})
    assert response.status_code == 200
    assert response.headers["cache-control"] == "private, no-store"
    exported = client.post("/profile-builder/export/docx", headers={"X-Profile-Builder-Access-Token": PROFILE_TOKEN}, json={
        "profile": candidate().model_dump(mode="json"), "anonymization": {},
        "template_id": "idego-default", "template": default_profile_template().model_dump(mode="json"),
    })
    assert exported.status_code == 200
    assert exported.headers["cache-control"] == "private, no-store"
    assert exported.headers["x-content-type-options"] == "nosniff"
